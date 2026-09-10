"""
etl/build_working_days_cache.py
================================
Builds cache/working_days.data.js -- the source for the "Field Working Days
Intelligence" tab (js/working-days.js). Separate from cache/sprint.json on
purpose: Sprint's cache only ever holds ONE live period (EVAL_PERIOD_NAME) at
a time, with past months archived to cache/sprint_history/ -- but this page's
whole point is a Feb-Jul trend view sourced straight from the KPI template's
already-stacked months, independent of whichever single month Sprint is
currently ranking. Re-run this any time Sprint_Missing_KPI_Template.xlsx's
DM_DSM/ASM/NSM tabs change (new month appended, TOT formula edited, etc.).

Reads:
  - zeta sprint/Sprint_Missing_KPI_Template.xlsx (DM_DSM, ASM, NSM tabs --
    Brand_Manager has no working-days data and is skipped, per Ahmed
    2026-09-05 "except bm").
  - cache/sprint.json, for the canonical BU/Line per employee code (the same
    majority-vote assignment already used to score Zeta Sprint -- far more
    reliable than the template's free-text Team column, see
    [[sprint_dmdsm_fielddays_added]] in project memory for why).
  - Database Shortcut.xlsx (2026-09-06, Ahmed "flage with hiring date and
    resignation date and if emp left the coumpany in filterd month flag as
    left company") -- the same authoritative HR source
    etl/build_sprint_cache.py itself reads for hire date / Last Day of Work,
    keyed by Code (Sheet1, col 0=Code, col 2=Hire Date, col 20=Last Day of
    Work). Every employee-month row gets hireDate/lastDay attached, plus a
    computed leftCompany flag (true when lastDay falls inside that row's own
    calendar month -- i.e. they resigned during the very month being
    viewed).

Writes:
  - cache/working_days.json       (plain JSON, for inspection/debugging)
  - cache/working_days.data.js    (gzip+base64, window.WORKING_DAYS_CACHE =
                                    {b64Data:"..."}), same encoding
                                    convention as cache/sprint.data.js.

TOT (Time Out of Territory, Ahmed 2026-09-05) = the off-day categories
subtracted from Calendar Days before applying each tier's multiplier:
Weekends, Leave Days, Holidays, AV Confrance, Business Travel, Confrance,
Gathering Meeting (DM_DSM only -- no such column on NSM/ASM), Group Meeting
(RTD), Sales Meeting, Training (absent on ASM).

Formula (confirmed with Ahmed 2026-09-05, DM_DSM added at this session --
NSM/ASM formula dates to 2026-09-02, see [[sprint_kpi_template_nsm]]):
  Target Working Days = (Calendar Days - TOT) * multiplier
    multiplier: DM_DSM 0.8, ASM 0.6, NSM 0.3
  Field Working Days % = All Visit Days / Target Working Days

This script does NOT write the Target Working Days / Field Working Days
columns back into the xlsx -- those are already there (written once by hand
this session; see [[sprint_dmdsm_fielddays_added]] and
[[sprint_kpi_template_nsm]]). It only READS them, same as
etl/build_sprint_cache.py's load_kpi_template() does for scoring.
"""
import openpyxl
import json
import re
import os
import datetime
import calendar
import base64
import gzip

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(BASE_DIR, 'zeta sprint', 'Sprint_Missing_KPI_Template.xlsx')
SPRINT_JSON_PATH = os.path.join(BASE_DIR, 'cache', 'sprint.json')
DB_PATH = os.path.join(BASE_DIR, 'Database Shortcut.xlsx')
OUT_JSON = os.path.join(BASE_DIR, 'cache', 'working_days.json')
OUT_JS = os.path.join(BASE_DIR, 'cache', 'working_days.data.js')

# schemaVersion 2 (2026-09-06): added hireDate/lastDay/leftCompany per
# employee-month row, sourced from Database Shortcut.xlsx.
SCHEMA_VERSION = 2

# Calendar-year bounds for each row's own month, used only to test whether
# a resignation (Last Day of Work) fell inside that specific month --
# same 2026 assumption the rest of this pipeline makes (EVAL_PERIOD_NAME
# etc. in etl/build_sprint_cache.py).
YEAR = 2026
MONTH_NUM = {'February': 2, 'March': 3, 'April': 4, 'May': 5, 'June': 6, 'July': 7}

DEDUCT_LABELS = ['Weekends', 'Leave Days', 'Holidays', 'AV Confrance', 'Business Travel',
                 'Confrance', 'Gathering Meeting', 'Group Meeting (RTD)', 'Sales Meeting', 'Training']
MULTIPLIER = {'DM_DSM': 0.8, 'ASM': 0.6, 'NSM': 0.3}
TARGET_HEADER = 'Target Working Days (Calculated)'
FIELDDAYS_HEADER = 'Field Working Days -- Actual % Achieved (e.g. 0.85 for 85%)'
MONTH_ORDER = ['February', 'March', 'April', 'May', 'June', 'July']
MONTH_ALIAS = {'Feb': 'February'}
SHEET_TO_TIERKEY = {'DM_DSM': 'dmDsm', 'ASM': 'asm', 'NSM': 'nsm'}


def log(msg):
    print(f'[{datetime.datetime.now().strftime("%H:%M:%S")}] {msg}')


def month_from_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%B')
    if v is not None:
        parts = re.split(r'[/\-.]', str(v).strip())
        if len(parts) == 3:
            try:
                _d, m, _y = (int(p) for p in parts)
                return datetime.date(2000, m, 1).strftime('%B')
            except Exception:
                return None
    return None


def month_bounds(month_name):
    """(first_day, last_day) date objects for `month_name` in YEAR."""
    n = MONTH_NUM[month_name]
    first = datetime.date(YEAR, n, 1)
    last = datetime.date(YEAR, n, calendar.monthrange(YEAR, n)[1])
    return first, last


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def main():
    log(f'reading {os.path.basename(TEMPLATE_PATH)} ...')
    wb = openpyxl.load_workbook(TEMPLATE_PATH, data_only=True)

    log('reading canonical BU/Line from cache/sprint.json ...')
    with open(SPRINT_JSON_PATH) as f:
        sp = json.load(f)

    log('reading hire/resignation dates from Database Shortcut.xlsx ...')
    wb_db = openpyxl.load_workbook(DB_PATH, data_only=True, read_only=True)
    ws_db = wb_db['Sheet1']
    code_to_hire, code_to_lastday = {}, {}
    for r in ws_db.iter_rows(min_row=2, values_only=True):
        code, hire, last_day = r[0], r[2], r[20]
        if code is None:
            continue
        key = str(int(code)) if isinstance(code, (int, float)) else str(code).strip()
        code_to_hire[key] = as_date(hire)
        code_to_lastday[key] = as_date(last_day)
    bu_line_by_sheet_code = {}
    for sheet_name, tier_key in SHEET_TO_TIERKEY.items():
        m = {}
        for rec in sp[tier_key]['ranked']:
            m[str(rec['code'])] = dict(bu=rec.get('bu'), line=rec.get('line'))
        # also fold in `excluded` (probation-not-passed, resigned, etc.) --
        # most of these still carry a real bu/line and were showing as
        # 'Unassigned' on the working-days leaderboard for no reason; a
        # resigned employee with bu=None/line=None correctly stays
        # Unassigned either way.
        for rec in sp[tier_key].get('excluded', []):
            code_s = str(rec['code'])
            if code_s not in m:
                m[code_s] = dict(bu=rec.get('bu'), line=rec.get('line'))
        bu_line_by_sheet_code[sheet_name] = m

    result = {}
    for sheet_name in ('DM_DSM', 'ASM', 'NSM'):
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

        def hidx(name):
            return header.index(name) if name in header else None

        code_i = hidx('Employee Code')
        code_i = code_i if code_i is not None else hidx('Code')
        name_i = hidx('Employee Name')
        name_i = name_i if name_i is not None else hidx('Name')
        profile_i = hidx('Profiles')
        date_i = hidx('Date')
        date_i = date_i if date_i is not None else hidx('DATE')
        month_i = hidx('Month')
        cal_i = hidx('Calendar Days')
        allvisit_i = hidx('All Visit Days')
        target_i = hidx(TARGET_HEADER)
        fd_i = hidx(FIELDDAYS_HEADER)
        ded_idx = {lbl: hidx(lbl) for lbl in DEDUCT_LABELS if hidx(lbl) is not None}

        rows_by_month = {}
        bu_line_map = bu_line_by_sheet_code[sheet_name]
        n_rows = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = row[code_i] if code_i is not None else None
            if code is None:
                continue
            if month_i is not None:
                m = row[month_i]
                m = MONTH_ALIAS.get(m, m)
            elif date_i is not None:
                m = month_from_date(row[date_i])
            else:
                m = None
            if m not in MONTH_ORDER:
                continue
            code_s = str(int(code)) if isinstance(code, (int, float)) else str(code).strip()
            name = row[name_i] if name_i is not None else None
            profile = row[profile_i] if profile_i is not None else None
            target = row[target_i] if target_i is not None else None
            fd = row[fd_i] if fd_i is not None else None
            allvisit = row[allvisit_i] if allvisit_i is not None else None
            cal = row[cal_i] if cal_i is not None else None
            deducts = {}
            for lbl, idx in ded_idx.items():
                v = row[idx]
                deducts[lbl] = float(v) if isinstance(v, (int, float)) else 0.0
            bl = bu_line_map.get(code_s, {})
            rows_by_month.setdefault(m, []).append(dict(
                code=code_s, name=name, profile=profile,
                bu=bl.get('bu'), line=bl.get('line'),
                calendarDays=cal, targetDays=target, fieldPct=fd,
                allVisitDays=allvisit, deducts=deducts,
                hireDate=code_to_hire.get(code_s), lastDay=code_to_lastday.get(code_s)))
            n_rows += 1
        result[sheet_name] = rows_by_month
        log(f'  {sheet_name}: {n_rows} rows across {len(rows_by_month)} month(s)')

    out = {
        'meta': {
            'schemaVersion': SCHEMA_VERSION,
            'generatedAt': datetime.datetime.now().isoformat(),
            'sourceFile': 'zeta sprint/Sprint_Missing_KPI_Template.xlsx',
        },
        'monthOrder': MONTH_ORDER,
        'multiplier': MULTIPLIER,
        'tiers': {},
        'employees': {},
    }
    for sheet_name in ('DM_DSM', 'ASM', 'NSM'):
        tier_out = {'months': {}}
        emp_by_month = {}
        for m in MONTH_ORDER:
            rows = result[sheet_name].get(m, [])
            if not rows:
                continue
            valid_fd = [r['fieldPct'] for r in rows if isinstance(r['fieldPct'], (int, float))]
            avg_fd = sum(valid_fd) / len(valid_fd) if valid_fd else None
            ded_totals = {}
            for r in rows:
                for lbl, v in r['deducts'].items():
                    ded_totals[lbl] = ded_totals.get(lbl, 0.0) + v
            n = len(rows)
            ded_avg = {lbl: round(v / n, 2) for lbl, v in ded_totals.items()}
            bands = {'below70': 0, '70to85': 0, '85to100': 0, 'above100': 0}
            for v in valid_fd:
                if v < 0.70:
                    bands['below70'] += 1
                elif v < 0.85:
                    bands['70to85'] += 1
                elif v <= 1.00:
                    bands['85to100'] += 1
                else:
                    bands['above100'] += 1
            tier_out['months'][m] = dict(count=n, avgFieldPct=round(avg_fd, 4) if avg_fd is not None else None,
                                          deductAvg=ded_avg, bands=bands)
            month_start, month_end = month_bounds(m)
            emp_rows = []
            for r in rows:
                last_day = r.get('lastDay')
                left_company = last_day is not None and month_start <= last_day <= month_end
                emp_rows.append(dict(
                    code=r['code'], name=r['name'],
                    bu=r['bu'] or 'Unassigned', line=r['line'] or 'Unassigned',
                    profile=r['profile'] if sheet_name == 'DM_DSM' else None,
                    calendarDays=r['calendarDays'],
                    deducts={k: round(v, 2) for k, v in r['deducts'].items()},
                    deductSum=round(sum(r['deducts'].values()), 2),
                    targetDays=round(r['targetDays'], 2) if isinstance(r['targetDays'], (int, float)) else None,
                    allVisitDays=r['allVisitDays'],
                    fieldPct=round(r['fieldPct'], 4) if isinstance(r['fieldPct'], (int, float)) else None,
                    hireDate=r['hireDate'].isoformat() if r.get('hireDate') else None,
                    lastDay=last_day.isoformat() if last_day else None,
                    leftCompany=left_company,
                ))
            emp_by_month[m] = emp_rows
        out['tiers'][sheet_name] = tier_out
        out['employees'][sheet_name] = emp_by_month

    with open(OUT_JSON, 'w') as f:
        json.dump(out, f)
    log(f'wrote {os.path.basename(OUT_JSON)}  {os.path.getsize(OUT_JSON) // 1024:,} KB')

    json_str = json.dumps(out)
    gz = gzip.compress(json_str.encode('utf-8'), compresslevel=9)
    b64 = base64.b64encode(gz).decode('ascii')
    with open(OUT_JS, 'w') as f:
        f.write('window.WORKING_DAYS_CACHE = {b64Data:"' + b64 + '"};\n')
    log(f'wrote {os.path.basename(OUT_JS)}  {os.path.getsize(OUT_JS) // 1024:,} KB (gzip+base64)')
    log('Field Working Days Intelligence cache complete.')


if __name__ == '__main__':
    main()
