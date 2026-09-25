#!/usr/bin/env python3
"""
build_list_intel_cache.py
=========================
ETL for the "List Intelligence" tab (CRM customer lists vs Promo Grid targets).

Replaces the frozen standalone page `List Intell/field_force_dashboard (9).html`,
whose data was typed into the HTML with no refresh path. This script rebuilds
the same data model from the source workbooks every cycle.

SOURCES  (folder: <project>/List Intell/)
-----------------------------------------------------------------------------
  All Lists*.xlsx          CRM export. Sheets "Total PM Lists", "Total AM Lists"
                           (one row per rep x plan) and "Details" (one row per
                           customer on a rep's list). Columns are located by
                           HEADER NAME, never by position.
  <LINE> Promo Grid.xlsx   One per line (16). The line key comes from the file
                           name: everything before " Promo"/" Prorm"/" Prormo"
                           (case-insensitive), upper-cased -> "Derma" = DERMA.
                           Cells are located by LABEL ("Doctors List", "AM List",
                           "Class A1", "Total", "No.Of MR", ...), never by fixed
                           address, because the 16 grids differ (sheet names,
                           row offsets, missing class rows, unlabeled rows,
                           "-   " text cells, numbers stored as text).
  The Pharmacies Plan block of each grid is out of scope and never read.

BUSINESS RULES  (confirmed by Ahmed, 2026-09-25)
-----------------------------------------------------------------------------
  * Target benchmark = the grid's FLAT per-rep list size ("Doctors List" /
    "AM List", rep side). The A1-C3 class breakdown is kept for display; any
    disagreement is surfaced as reconFlag + reconNotes, never "fixed".
  * Line target = flat target per rep x Planned MR ("No.Of MR" on the grid).
    Actual in-post reps and vacancies are carried per rep (vacant flag) so the
    UI can show capacity gaps next to it.
  * CRM capacity = CRM target call rate x WORKING_DAYS[plan]  (PM 20, AM 22).
    The CRM export's own Capacity column is always call rate x 20, so AM rows
    are re-based to 22 here; the source values are kept as *Source fields.
    Deviation = CRM frequency - capacity; status Under / Balanced / Over.
  * CHC_Sales is merged into CHC (lineRemapped + origLine kept for audit).
  * Customer rows are published WITHOUT street address (privacy, payload).
  * Tolerance band (+/-10%) is applied in the UI, not baked in here.

OUTPUT
-----------------------------------------------------------------------------
  cache/list_intel.json      plain JSON (inspection / validation)
  cache/list_intel.data.js   window.LIST_INTEL_CACHE = {b64Data:"<gzip+base64 JSON>"}
                             -- same convention as market_intel.data.js; loaded
                             lazily by js/cache-loader.js when the tab opens.

Run:  python etl/build_list_intel_cache.py            (from the project root)
Exit code 0 = built; 1 = a source problem that makes the output unsafe.
"""
import os, sys, re, glob, json, gzip, base64, hashlib, datetime
from collections import OrderedDict, Counter

import openpyxl

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_DIR = os.path.join(ROOT_DIR, 'List Intell')
OUT_JSON = os.path.join(ROOT_DIR, 'cache', 'list_intel.json')
OUT_JS = os.path.join(ROOT_DIR, 'cache', 'list_intel.data.js')

CLASSES = ['A1', 'A2', 'A3', 'B1', 'B2', 'B3', 'C1', 'C2', 'C3']
WORKING_DAYS = {'PM': 20, 'AM': 22}
LINE_REMAP = {'CHC_SALES': 'CHC'}
TOLERANCE_PCT = 10
PLAN_SHEETS = {'PM': 'Total PM Lists', 'AM': 'Total AM Lists'}
DETAILS_SHEET = 'Details'
CUSTOMER_COLUMNS = ['name', 'type', 'specialty', 'class', 'area', 'frequency', 'clinicGroup']

WARNINGS = []


def log(msg):
    print('[list_intel] ' + msg, flush=True)


def warn(msg):
    WARNINGS.append(msg)
    log('WARNING: ' + msg)


def fail(msg):
    log('ERROR: ' + msg)
    sys.exit(1)


# ---------------------------------------------------------------- cell cleaning
def norm(s):
    """Label normaliser: lower-case, collapse spaces/dots, strip."""
    if s is None:
        return ''
    return re.sub(r'[\s.]+', ' ', str(s)).strip().lower()


def num(v):
    """Tolerant numeric read. None / '' / '-' / '-   ' / '`' -> None;
    '30' -> 30; 24.5 stays 24.5; integral floats become int."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        t = str(v).strip().replace(',', '')
        if t in ('', '-', '`') or re.fullmatch(r'-+\s*', t):
            return None
        try:
            f = float(t)
        except ValueError:
            return None
    return int(f) if f == int(f) else round(f, 4)


def txt(v):
    """Text cell as stored; blank ('' / spaces, which the read-only reader
    returns for empty shared strings) -> None."""
    if v is None or (isinstance(v, str) and v.strip() == ''):
        return None
    return v


def clean_name(v):
    return '' if v is None else str(v).strip()


# ---------------------------------------------------------------- promo grids
GRID_NAME_RE = re.compile(r'^(.+?)\s+pro[a-z]*\s*grid', re.I)


def discover_grids():
    files = [f for f in glob.glob(os.path.join(SOURCE_DIR, '*.xlsx'))
             if 'grid' in os.path.basename(f).lower() and not os.path.basename(f).startswith('~$')]
    by_line = {}
    for f in sorted(files):
        m = GRID_NAME_RE.match(os.path.basename(f))
        if not m:
            warn('promo grid file name not understood, skipped: ' + os.path.basename(f))
            continue
        line = m.group(1).strip().upper()
        if line in by_line:
            a, b = by_line[line], f
            if hashlib.md5(open(a, 'rb').read()).hexdigest() == hashlib.md5(open(b, 'rb').read()).hexdigest():
                warn('duplicate promo grid for %s (identical), using %s' % (line, os.path.basename(a)))
                continue
            fail('two DIFFERENT promo grids for line %s: %s / %s -- keep one' %
                 (line, os.path.basename(a), os.path.basename(b)))
        by_line[line] = f
    return by_line


def grid_sheet(wb):
    """The class-plan sheet: the first sheet holding both 'Physicians Plan' and 'Doctors List'."""
    for ws in wb.worksheets:
        labels = set()
        for row in ws.iter_rows(min_row=1, max_row=40, max_col=3, values_only=True):
            for v in row:
                labels.add(norm(v))
        if 'doctors list' in labels and any(l.startswith('physicians plan') for l in labels):
            return ws
    return None


def parse_block(grid, r0, r_end, plan_label):
    """Parse one plan block (rows r0..r_end inclusive) of the rep side (cols A-H)
    and line side (cols J-Q). `grid` is {(row, col_letter): value}."""
    def g(r, c):
        return grid.get((r, c))

    e = OrderedDict()
    e['targetPerRep'] = num(g(r0, 'B'))
    rep_meta, line_meta = {}, {}
    cls_t, cls_f = {c: None for c in CLASSES}, {c: None for c in CLASSES}
    seen_class = False
    total_rep = total_line = None
    for r in range(r0, r_end + 1):
        la = norm(g(r, 'A'))
        if la:
            rep_meta[la] = g(r, 'B')
        lj = norm(g(r, 'J'))
        if lj:
            line_meta[lj] = g(r, 'K')
        le = norm(g(r, 'E'))
        m = re.fullmatch(r'class ([abc][123])', le)
        if m:
            k = m.group(1).upper()
            seen_class = True
            cls_t[k] = num(g(r, 'F'))
            cls_f[k] = num(g(r, 'G'))
        elif le == 'total':
            total_rep, total_line = num(g(r, 'F')), num(g(r, 'O'))
    e['_workingDaysGrid'] = num(rep_meta.get('working days'))
    e['_visitsPerDay'] = num(rep_meta.get('visits / day'))
    e['_visitsPerCycle'] = num(rep_meta.get('visits / cycle'))
    e['_lineTargetStated'] = num(line_meta.get(plan_label))
    mr = None
    for k, v in line_meta.items():
        if re.fullmatch(r'no ?of mr', k.replace('.', '')):
            mr = num(v)
    e['_mr'] = mr
    e['_classesTarget'] = cls_t
    e['_classesFreq'] = cls_f
    e['_seenClass'] = seen_class
    e['_totalRep'] = total_rep
    e['_totalLine'] = total_line
    return e


def parse_grid(line, path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = grid_sheet(wb)
    if ws is None:
        fail('%s: no sheet with "Physicians Plan" + "Doctors List"' % os.path.basename(path))
    grid = {}
    anchors = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 120)):
        for c in row:
            if c.value is not None and c.value != '':
                grid[(c.row, c.column_letter)] = c.value
        a = norm(row[0].value) if row else ''
        if a == 'doctors list' and 'PM' not in anchors:
            anchors['PM'] = row[0].row
        elif a == 'am list' and 'AM' not in anchors:
            anchors['AM'] = row[0].row
        elif a == 'pharmacy list' and 'PH' not in anchors:
            anchors['PH'] = row[0].row
    if 'PM' not in anchors or 'AM' not in anchors:
        fail('%s: could not find the "Doctors List" / "AM List" blocks' % os.path.basename(path))
    pm_end = anchors['AM'] - 1
    am_end = (anchors.get('PH') or (anchors['AM'] + 14)) - 1
    pm = parse_block(grid, anchors['PM'], pm_end, 'doctors list')
    am = parse_block(grid, anchors['AM'], am_end, 'am list')
    mr = pm['_mr'] if pm['_mr'] is not None else am['_mr']
    if mr is None:
        fail('%s: Planned MR ("No.Of MR") not found' % os.path.basename(path))
    out = {}
    for plan, blk in (('PM', pm), ('AM', am)):
        out[plan] = finalize_target(line, plan, blk, mr, ws.title, os.path.basename(path))
    return out


def finalize_target(line, plan, b, mr, sheet, fname):
    t = b['targetPerRep']
    if t is None:
        fail('%s %s: flat target per rep is empty (%s)' % (line, plan, fname))
    cls_t, cls_f = b['_classesTarget'], b['_classesFreq']
    values = [v for v in cls_t.values() if v is not None]
    no_seg = (not b['_seenClass']) or not any(values)
    sum_rep = sum(values) if values else 0
    sum_rep = int(sum_rep) if sum_rep == int(sum_rep) else sum_rep
    e = OrderedDict()
    e['targetPerRep'] = t
    e['plannedMR'] = mr
    e['classesTarget'] = None if no_seg else cls_t
    e['classesFreqTarget'] = None if no_seg else cls_f
    e['noSegmentation'] = no_seg
    e['sumClassesPerRep'] = 0 if no_seg else sum_rep
    e['segTotalRowPerRep'] = 0 if no_seg else b['_totalRep']
    e['lineTargetStated'] = b['_lineTargetStated'] if b['_lineTargetStated'] is not None else t * mr
    e['sumClassesLine'] = 0 if no_seg else num(sum_rep * mr)
    e['segTotalRowLine'] = 0 if no_seg else b['_totalLine']
    notes = []
    if not no_seg:
        if t != sum_rep:
            notes.append('Flat target/rep (%s) != sum of A1-C3/rep (%s)' % (fmt(t), fmt(sum_rep)))
        if e['segTotalRowPerRep'] is None:
            notes.append('Segmentation Total row (per rep) not entered in source')
        elif t != e['segTotalRowPerRep']:
            notes.append('Flat target/rep (%s) != segmentation Total row/rep (%s)' % (fmt(t), fmt(e['segTotalRowPerRep'])))
        per_rep_mr = t * mr
        if e['segTotalRowLine'] is not None and per_rep_mr != e['segTotalRowLine']:
            notes.append('Per-rep x MR (%s) != Line segmentation Total row (%s)' % (fmt(per_rep_mr), fmt(e['segTotalRowLine'])))
    if b['_lineTargetStated'] is not None and t * mr != b['_lineTargetStated']:
        notes.append('Per-rep x MR (%s) != sheet Line-total cell (%s)' % (fmt(t * mr), fmt(b['_lineTargetStated'])))
    e['reconFlag'] = bool(notes)
    e['reconNotes'] = notes
    e['targetVisitsPerDay'] = b['_visitsPerDay']
    e['targetCallsPerCycle'] = b['_visitsPerCycle']
    e['workingDaysGrid'] = b['_workingDaysGrid']
    e['source'] = {'file': fname, 'sheet': sheet}
    return e


def fmt(v):
    return str(int(v)) if isinstance(v, float) and v == int(v) else str(v)


# ---------------------------------------------------------------- CRM lists
def find_lists_workbook():
    c = [f for f in glob.glob(os.path.join(SOURCE_DIR, 'All Lists*.xlsx')) if not os.path.basename(f).startswith('~$')]
    if not c:
        fail('no "All Lists*.xlsx" in ' + SOURCE_DIR)
    if len(c) > 1:
        c.sort(key=os.path.getmtime, reverse=True)
        warn('several lists workbooks, using the newest: ' + os.path.basename(c[0]))
    return c[0]


def header_index(header, wanted, sheet):
    h = [norm(x) for x in header]
    out = {}
    for key, options in wanted.items():
        idx = next((h.index(norm(o)) for o in options if norm(o) in h), None)
        if idx is None:
            fail('sheet "%s": column for %s not found (looked for %s)' % (sheet, key, options))
        out[key] = idx
    return out


PLAN_COLS = {
    'bu': ['BU Name'], 'line': ['Line'], 'area': ['Area'], 'manager': ['Manager'], 'title': ['Title'],
    'employee': ['Employee'],
    'current': ['Total Number Of Drs', 'Total Number AM Accounts', 'Total Number Of AM Accounts'],
    'freq': ['Total Number Of Frequency', 'Frequency'],
    'rate': ['Target Call Rate PM', 'Target Call Rate AM', 'Target Call Rate'],
    'capacity': ['Capacity'], 'deviation': ['Deviation'], 'status': ['Status Of Deviation'],
}
for _c in CLASSES + ['Others']:
    PLAN_COLS['cls_' + _c] = ['Total Number Of Class ' + _c]


def status_for(dev):
    return 'Balanced' if dev == 0 else ('Under Capacity' if dev < 0 else 'Over Capacity')


def read_reps(wb, promo):
    reps = []
    for plan, sheet in PLAN_SHEETS.items():
        if sheet not in wb.sheetnames:
            fail('lists workbook has no sheet "%s"' % sheet)
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows)
        ix = header_index(header, PLAN_COLS, sheet)
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            emp = clean_name(row[ix['employee']])
            orig_line = clean_name(row[ix['line']])
            if not emp and not orig_line:
                continue
            line = LINE_REMAP.get(orig_line.upper(), orig_line.upper())
            if line not in promo:
                warn('%s: line "%s" has no promo grid -- rep %s skipped' % (plan, orig_line, emp))
                continue
            pt = promo[line][plan]
            rate = num(row[ix['rate']])
            cap_src = num(row[ix['capacity']])
            freq = num(row[ix['freq']])
            cap = None if rate is None else rate * WORKING_DAYS[plan]
            dev = None if (cap is None or freq is None) else freq - cap
            r = OrderedDict()
            r['bu'] = clean_name(row[ix['bu']])
            r['line'] = line
            r['origLine'] = orig_line
            r['lineRemapped'] = line != orig_line
            r['area'] = clean_name(row[ix['area']])
            r['manager'] = clean_name(row[ix['manager']])
            r['employee'] = emp
            r['title'] = clean_name(row[ix['title']])
            r['plan'] = plan
            r['vacant'] = 'vacant' in emp.lower()
            r['current'] = num(row[ix['current']]) or 0
            r['classesCurrent'] = OrderedDict((c, num(row[ix['cls_' + c]]) or 0) for c in CLASSES + ['Others'])
            r['target'] = pt['targetPerRep']
            r['plannedMR'] = pt['plannedMR']
            r['classesTarget'] = pt['classesTarget']
            r['classesFreqTarget'] = pt['classesFreqTarget']
            r['noSegTarget'] = pt['noSegmentation']
            r['reconFlag'] = pt['reconFlag']
            r['crmFreq'] = freq
            r['crmTargetCallRate'] = rate
            r['crmWorkingDays'] = WORKING_DAYS[plan]
            r['crmCapacity'] = cap
            r['crmDeviation'] = dev
            r['crmStatus'] = None if dev is None else status_for(dev)
            r['crmCapacitySource'] = cap_src
            r['crmDeviationSource'] = num(row[ix['deviation']])
            r['crmStatusSource'] = clean_name(row[ix['status']])
            r['targetCallsPerCycle'] = pt['targetCallsPerCycle']
            r['targetVisitsPerDay'] = pt['targetVisitsPerDay']
            reps.append(r)
    return reps


DETAIL_COLS = {
    'employee': ['Employee'], 'type': ['Type'], 'name': ['Customer Name'], 'specialty': ['Specialty'],
    'class': ['Class'], 'area': ['Area'], 'clinic': ['Clinic Group'], 'freq': ['Frequency'],
}


def read_customers(wb):
    if DETAILS_SHEET not in wb.sheetnames:
        fail('lists workbook has no sheet "%s"' % DETAILS_SHEET)
    rows = wb[DETAILS_SHEET].iter_rows(values_only=True)
    ix = header_index(next(rows), DETAIL_COLS, DETAILS_SHEET)
    cust = OrderedDict()
    n = 0
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        emp = clean_name(row[ix['employee']])
        if not emp:
            continue
        # Street address is deliberately NOT read (privacy decision 2026-09-25).
        cust.setdefault(emp, []).append([
            txt(row[ix['name']]), txt(row[ix['type']]), txt(row[ix['specialty']]), txt(row[ix['class']]),
            txt(row[ix['area']]), num(row[ix['freq']]), txt(row[ix['clinic']]),
        ])
        n += 1
    return cust, n


# ---------------------------------------------------------------- checks
def sanity(reps, customers, promo):
    emp_rows = Counter((r['employee'], r['plan']) for r in reps)
    dups = [k for k, n in emp_rows.items() if n > 1 and 'vacant' not in k[0].lower()]
    if dups:
        warn('%d employee/plan pairs appear more than once, e.g. %s' % (len(dups), dups[:3]))
    active = {r['employee'] for r in reps if not r['vacant']}
    missing = sorted(e for e in active if e not in customers)
    if missing:
        warn('%d active reps have no customer rows in Details, e.g. %s' % (len(missing), missing[:3]))
    orphan = sorted(e for e in customers if e not in {r['employee'] for r in reps})
    if orphan:
        warn('%d Details employees are not in the PM/AM list sheets, e.g. %s' % (len(orphan), orphan[:3]))
    for r in reps:
        if r['crmCapacitySource'] is not None and r['crmTargetCallRate'] is not None \
                and r['crmCapacitySource'] != r['crmTargetCallRate'] * 20:
            warn('CRM Capacity is not call rate x 20 for %s %s' % (r['employee'], r['plan']))
            break


# ---------------------------------------------------------------- main
def main():
    started = datetime.datetime.now()
    if not os.path.isdir(SOURCE_DIR):
        fail('source folder not found: ' + SOURCE_DIR)
    grids = discover_grids()
    if not grids:
        fail('no promo grid workbooks found in ' + SOURCE_DIR)
    promo = OrderedDict()
    for line, path in grids.items():
        promo[line] = parse_grid(line, path)
    log('promo grids: %d lines (%s)' % (len(promo), ', '.join(promo)))

    lists_path = find_lists_workbook()
    wb = openpyxl.load_workbook(lists_path, read_only=True, data_only=True)
    reps = read_reps(wb, promo)
    customers, n_cust = read_customers(wb)
    wb.close()
    log('reps: %d rows (%d PM, %d AM, %d vacant); customers: %d rows for %d employees' % (
        len(reps), sum(r['plan'] == 'PM' for r in reps), sum(r['plan'] == 'AM' for r in reps),
        sum(r['vacant'] for r in reps), n_cust, len(customers)))
    sanity(reps, customers, promo)

    data = OrderedDict()
    data['meta'] = OrderedDict([
        ('builtAt', started.strftime('%Y-%m-%d %H:%M')),
        ('sources', OrderedDict([
            ('lists', os.path.basename(lists_path)),
            ('listsModified', datetime.datetime.fromtimestamp(os.path.getmtime(lists_path)).strftime('%Y-%m-%d %H:%M')),
            ('grids', OrderedDict((l, os.path.basename(p)) for l, p in grids.items())),
        ])),
        ('rules', OrderedDict([
            ('targetBasis', 'flat per-rep list size from the Promo Grid'),
            ('lineTargetBasis', 'target per rep x Planned MR'),
            ('tolerancePct', TOLERANCE_PCT),
            ('toleranceDefaultOn', True),
            ('crmWorkingDays', WORKING_DAYS),
            ('lineRemap', LINE_REMAP),
            ('pharmacyPlan', 'excluded'),
            ('customerAddress', 'excluded'),
        ])),
        ('customerColumns', CUSTOMER_COLUMNS),
        ('warnings', WARNINGS),
    ])
    data['promoTargets'] = promo
    data['reps'] = reps
    data['customers'] = customers

    os.makedirs(os.path.dirname(OUT_JSON), exist_ok=True)
    js = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
    tmp = OUT_JSON + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        f.write(js)
    os.replace(tmp, OUT_JSON)
    b64 = base64.b64encode(gzip.compress(js.encode('utf-8'), compresslevel=9)).decode('ascii')
    tmp = OUT_JS + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        f.write('window.LIST_INTEL_CACHE = {b64Data:"' + b64 + '"};\n')
    os.replace(tmp, OUT_JS)
    log('wrote %s (%s KB) and %s (%s KB gzip+base64) in %.1fs, %d warning(s)' % (
        os.path.basename(OUT_JSON), format(os.path.getsize(OUT_JSON) // 1024, ','),
        os.path.basename(OUT_JS), format(os.path.getsize(OUT_JS) // 1024, ','),
        (datetime.datetime.now() - started).total_seconds(), len(WARNINGS)))


if __name__ == '__main__':
    main()
