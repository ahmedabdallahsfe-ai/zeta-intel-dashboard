"""
leave_rules.py
==============
THE single source of truth for the SFE Sick Leave Impact Rule.

Both ETLs import from here -- refresh.py (Coverage / Right Frequency) and
etl/build_coaching_cache.py (Coaching Intelligence DV Coverage). That is the
whole point of this file: the rule decides who is excluded from a rate, from a
ranking and from a coaching denominator, and this project has already been
bitten more than once by two engines computing "the same" thing slightly
differently and quietly disagreeing. There must be exactly one definition of
"how many leave days did this rep have that month" and exactly one definition
of which band that puts them in.

If you change a threshold here, you are changing it everywhere. That is
intended. Do not copy these functions into an ETL.

--------------------------------------------------------------------------
THE RULE (Ahmed, 2026-09-15)
--------------------------------------------------------------------------
Evaluated per rep, per MONTH, on that month's own days -- never cumulative
across months, and never carried forward:

    0 - 5 days   Normal    Evaluate on 100% standard targets.
    6 - 15 days  Moderate  Prorate the Right Frequency target by the
                           active-working-days ratio. Coverage reach is
                           NEVER prorated.
    > 15 days    Excluded  Flag the territory: out of every Coverage/RF rate
    or Maternity           and out of every competitive ranking. Also out of
                           the Coaching DV Coverage denominator -- a manager
                           cannot double-visit someone who was not there
                           (Ahmed, 2026-09-16).

Days are CALENDAR days, weekends included, exactly as the Sick Leave report's
"Total" column states them. A weekend-aware variant was built 2026-09-15 and
reverted the same day at Ahmed's instruction -- do not reintroduce it without
asking him first.

Only Type == "Sick" counts toward the day bands. Annual, Unpaid and Marriage
are tracked in the same report but do not inflate the count, because a planned
absence is not the same commercial fact as an unplanned one. (Open question
with Ahmed as of 2026-09-16 -- if he extends the rule to those types, change
SICK_TYPES below and nothing else.) Maternity is its own unconditional
exclusion regardless of day count.
"""

from __future__ import annotations

import datetime
import os

# --- Band thresholds -------------------------------------------------------
SICK_LEAVE_NORMAL_MAX_DAYS = 5.0      # 0-5   -> Normal
SICK_LEAVE_MODERATE_MAX_DAYS = 15.0   # 6-15  -> Moderate;  >15 -> Excluded
STANDARD_MONTHLY_WORKING_DAYS = 22.0  # divisor for the active-working-days ratio
TIER_A_CLASS_PREFIX = "A"             # Class values starting with this = Tier A

# Leave types that count toward the day bands. Maternity is handled
# separately as an unconditional exclusion, so it is deliberately not here.
SICK_TYPES = {"sick"}
MATERNITY_TYPES = {"maternity"}

LEAVE_REPORT_FILENAME = "Sick Leave report.xlsx"

BAND_NORMAL = "Normal"
BAND_MODERATE = "Moderate"
BAND_EXCLUDED = "Excluded"

CANONICAL_MONTH_NAMES = {
    "jan": "January", "feb": "February", "mar": "March", "apr": "April",
    "may": "May", "jun": "June", "jul": "July", "aug": "August",
    "sep": "September", "oct": "October", "nov": "November", "dec": "December",
}

# Coaching Intelligence keys its months "2026-07"; the Coverage ETL keys them
# "July". Both are first-class, so the mapping lives here rather than in
# whichever file happened to need it first.
MONTH_NUM_TO_NAME = {
    1: "January", 2: "February", 3: "March", 4: "April", 5: "May", 6: "June",
    7: "July", 8: "August", 9: "September", 10: "October", 11: "November",
    12: "December",
}


def month_key_to_name(month_key: str) -> str | None:
    """'2026-07' -> 'July'. Returns None for anything unparseable."""
    try:
        return MONTH_NUM_TO_NAME[int(str(month_key).split("-")[1])]
    except (IndexError, ValueError, KeyError):
        return None


def normalize_code(value) -> str:
    """Employee Code as a bare string ('1966'), however the sheet stored it."""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    if isinstance(value, (int, float)):
        return str(int(value))
    return str(value).strip()


def resolve_effective_leave_days(total_days: float, doctor_action) -> float:
    """The report's 'Total' column is the requested date span; some rows carry
    a free-text 'Doctor Action' note showing fewer days were actually
    approved. Ahmed confirmed 2026-09-15 the rule uses the approved count
    where a note overrides it. Mapping built from the 267-row report as it
    stood that day -- extend it if a new note text appears:
        "Rejected by Dr" (incl. "...and add to annual")  -> 0 days
        "Only 7 Days"                                     -> 7
        "1 Week"                                          -> 7
        "three days gerented[sic] from company"           -> 3
        anything else ("Deduction from Salary", "Policy Company",
        "still in leave", "Maternity Leave", or no note)  -> Total as-is;
        those describe how the leave was handled, not that fewer days were
        actually taken.
    """
    action = (str(doctor_action) if doctor_action is not None else "").strip().lower()
    if not action or action in ("nan", "none"):
        return total_days
    if action.startswith("rejected"):
        return 0.0
    if "only 7 days" in action:
        return 7.0
    if action == "1 week":
        return 7.0
    if "three days" in action:
        return 3.0
    return total_days


def split_days_by_month(date_from, effective_days: float) -> dict[str, float]:
    """Walk `effective_days` calendar days from date_from and return
    {CanonicalMonthName: days}. A leave starting near month-end spills into
    the next month and is split across both -- the original implementation
    (pd.date_range(..., freq='MS')) silently dropped that spillover and
    attributed the whole block to one month.

    Days are CALENDAR days, weekends included -- see this module's header.
    """
    if effective_days is None or effective_days <= 0 or date_from is None:
        return {}
    counts: dict[str, float] = {}
    d = date_from
    remaining = float(effective_days)
    while remaining > 0:
        take = min(1.0, remaining)
        mname = CANONICAL_MONTH_NAMES.get(d.strftime("%b").lower(), d.strftime("%B"))
        counts[mname] = counts.get(mname, 0.0) + take
        remaining -= take
        d = d + datetime.timedelta(days=1)
    return counts


def leave_band(sick_days: float, is_maternity: bool = False) -> str:
    """The band a rep falls in for ONE month. Maternity is unconditional."""
    if is_maternity:
        return BAND_EXCLUDED
    days = float(sick_days or 0)
    if days > SICK_LEAVE_MODERATE_MAX_DAYS:
        return BAND_EXCLUDED
    if days > SICK_LEAVE_NORMAL_MAX_DAYS:
        return BAND_MODERATE
    return BAND_NORMAL


def active_ratio(sick_days: float) -> float:
    """Share of the standard working month the rep was actually available."""
    days = float(sick_days or 0)
    return max(0.0, (STANDARD_MONTHLY_WORKING_DAYS - days) / STANDARD_MONTHLY_WORKING_DAYS)


def _coerce_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    return None


def load_leave_report(path: str) -> dict:
    """Read the Sick Leave report once and return plain dicts both ETLs can
    use without pandas:

        {
          "found": bool,
          "path": str,
          "sickDays":   {(code, "July"): days, ...},   # Type=Sick only
          "maternity":  {(code, "July"), ...},         # month-scoped
          "sickDaysUndated": {code: days},   # rows with no Date from
          "maternityUndated": {code, ...},
          "rowsRead": int,
          "typeCounts": {"sick": n, ...},
        }

    Uses openpyxl so it stays importable from either ETL (Coverage's runs on
    pandas, Coaching's on openpyxl).
    """
    out = {
        "found": False, "path": path,
        "sickDays": {}, "maternity": set(),
        "sickDaysUndated": {}, "maternityUndated": set(),
        "rowsRead": 0, "typeCounts": {},
    }
    if not path or not os.path.exists(path):
        return out

    import openpyxl  # imported lazily so importing this module stays cheap

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        return out
    idx = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}

    def col(*names):
        for n in names:
            if n in idx:
                return idx[n]
        return None

    c_code = col("code")
    c_type = col("type")
    c_from = col("date from")
    c_total = col("total")
    c_action = col("doctor action")
    if c_code is None or c_total is None:
        wb.close()
        return out

    for row in rows:
        if row is None or c_code >= len(row):
            continue
        code = normalize_code(row[c_code])
        if not code:
            continue
        out["rowsRead"] += 1
        # Case-normalised: the source carries both "Marriage" and "marriage",
        # and a case-sensitive match would silently drop half of them.
        ltype = (str(row[c_type]).strip().lower() if c_type is not None and row[c_type] is not None else "")
        out["typeCounts"][ltype] = out["typeCounts"].get(ltype, 0) + 1

        raw_total = row[c_total] if c_total < len(row) else 0
        try:
            raw_total = float(raw_total or 0)
        except (TypeError, ValueError):
            raw_total = 0.0
        action = row[c_action] if (c_action is not None and c_action < len(row)) else None
        eff = resolve_effective_leave_days(raw_total, action)
        dfrom = _coerce_date(row[c_from]) if (c_from is not None and c_from < len(row)) else None

        is_sick = ltype in SICK_TYPES
        is_mat = ltype in MATERNITY_TYPES
        if not (is_sick or is_mat):
            continue

        if dfrom is None:
            if is_sick and eff > 0:
                out["sickDaysUndated"][code] = out["sickDaysUndated"].get(code, 0.0) + eff
            if is_mat:
                out["maternityUndated"].add(code)
            continue

        months = split_days_by_month(dfrom, eff) if eff > 0 else {}
        if is_mat and not months:
            # A maternity row with an overridden zero still marks the month.
            mname = CANONICAL_MONTH_NAMES.get(dfrom.strftime("%b").lower(), dfrom.strftime("%B"))
            months = {mname: 0.0}
        for mname, days in months.items():
            if is_sick:
                out["sickDays"][(code, mname)] = out["sickDays"].get((code, mname), 0.0) + days
            if is_mat:
                out["maternity"].add((code, mname))

    wb.close()
    out["found"] = True
    return out


def load_leave_rows(path: str) -> dict:
    """Per-employee RAW leave rows, for UI detail popups.

    load_leave_report() above answers "which band is this rep in this month"
    and deliberately keeps only what the rule needs -- Sick day counts and a
    Maternity flag. That is the right shape for scoring and the wrong shape
    for showing a manager WHY a rep's KPI collapsed. Zeta Sprint (Ahmed,
    2026-09-16: "if zero in kpi coverage and right frequency due to leave
    flag leave type and popup duration") needs the leave TYPE and the actual
    date span, which the band math throws away.

    So this returns every leave row as recorded, ALL types included --
    Annual/Unpaid/Marriage too, even though they do not count toward the day
    bands. A manager looking at a zero should see the whole month's absence
    picture, with the band-driving Sick days called out separately rather
    than the non-counting types silently hidden (which reads as missing
    data). `countsTowardBand` marks which is which, so no consumer has to
    re-derive SICK_TYPES for itself.

        {code: [{"type": "Sick", "dateFrom": "2026-07-03", "dateTo": ...,
                 "totalDays": 15.0, "effectiveDays": 7.0,
                 "doctorAction": "Only 7 Days", "months": {"July": 7.0},
                 "countsTowardBand": True, "isMaternity": False}, ...]}

    effectiveDays is the Doctor-Action-resolved count the rule actually uses
    (see resolve_effective_leave_days) -- shown alongside totalDays so an
    overridden row explains itself instead of looking like a data error.
    """
    out: dict[str, list] = {}
    if not path or not os.path.exists(path):
        return out

    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        return out
    idx = {str(h).strip().lower(): i for i, h in enumerate(header) if h is not None}

    def col(*names):
        for n in names:
            if n in idx:
                return idx[n]
        return None

    # "day" appears TWICE in this header (weekday name after each date), so
    # a plain dict comprehension keeps only the second. Both are ignored
    # here -- only Date from / Date to are read.
    c_code, c_type = col("code"), col("type")
    c_from, c_to = col("date from"), col("date to")
    c_total, c_action = col("total"), col("doctor action")
    if c_code is None or c_total is None:
        wb.close()
        return out

    for row in rows:
        if row is None or c_code >= len(row):
            continue
        code = normalize_code(row[c_code])
        if not code:
            continue
        raw_type = (str(row[c_type]).strip() if c_type is not None and row[c_type] is not None else "")
        ltype = raw_type.lower()
        try:
            total = float((row[c_total] if c_total < len(row) else 0) or 0)
        except (TypeError, ValueError):
            total = 0.0
        action = row[c_action] if (c_action is not None and c_action < len(row)) else None
        action_text = str(action).strip() if action is not None and str(action).strip().lower() not in ("nan", "none") else ""
        eff = resolve_effective_leave_days(total, action)
        dfrom = _coerce_date(row[c_from]) if (c_from is not None and c_from < len(row)) else None
        dto = _coerce_date(row[c_to]) if (c_to is not None and c_to < len(row)) else None
        months = split_days_by_month(dfrom, eff) if (dfrom and eff > 0) else {}
        if not months and dfrom is not None:
            mname = CANONICAL_MONTH_NAMES.get(dfrom.strftime("%b").lower(), dfrom.strftime("%B"))
            months = {mname: 0.0}
        out.setdefault(code, []).append({
            "type": raw_type or "Leave",
            "dateFrom": dfrom.isoformat() if dfrom else None,
            "dateTo": dto.isoformat() if dto else None,
            "totalDays": total,
            "effectiveDays": eff,
            "doctorAction": action_text,
            "months": months,
            "countsTowardBand": ltype in SICK_TYPES,
            "isMaternity": ltype in MATERNITY_TYPES,
        })

    wb.close()
    return out


def rows_in_month(rows: list, month_name: str) -> list:
    """Filter load_leave_rows()' per-code list down to one month, annotating
    each surviving row with monthDays -- how many of its days landed in THAT
    month. A leave spanning a month boundary therefore shows its true
    in-month weight (e.g. a 28-day block starting 7 Jan contributes 25 days
    to January and 3 to February), not its full span."""
    out = []
    for r in rows or []:
        days = (r.get("months") or {}).get(month_name)
        if days is None:
            continue
        rr = dict(r)
        rr["monthDays"] = days
        out.append(rr)
    return out


def band_for(leave: dict, code: str, month_name: str) -> str:
    """Convenience: the band for one rep in one month, given a loaded report."""
    code = normalize_code(code)
    days = leave["sickDays"].get((code, month_name), leave["sickDaysUndated"].get(code, 0.0))
    is_mat = (code, month_name) in leave["maternity"] or code in leave["maternityUndated"]
    return leave_band(days, is_mat)


def is_absent_for_coaching(leave: dict, code: str, month_name: str) -> bool:
    """Should this rep be dropped from a manager's DV Coverage denominator
    for this month? (Ahmed, 2026-09-16: "if one team member absent or take
    leave for full month ... his dv coverage is 100%".)

    Uses the SAME Excluded band the Coverage rule uses, deliberately -- one
    definition of "not available this month" across the whole platform. A
    second, coaching-only threshold is exactly how the two engines drift.
    """
    return band_for(leave, code, month_name) == BAND_EXCLUDED
